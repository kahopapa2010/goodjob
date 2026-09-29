"""マスタ（M01〜M09）・月初在庫（T90）・検証結果ブックを作る
使い方: python build_masters.py <月次日報.xlsx> <CSVフォルダ（extract.pyの出力先）>
extract.py を先に実行しておくこと（X_検証_*.csv と T*.csv を使う）。
"""
import sys, os, csv, glob, datetime, warnings
import openpyxl
from xlsx_load import load
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as gl
warnings.filterwarnings("ignore")

src, CSVD = sys.argv[1], sys.argv[2]
wv = load(src, data_only=True)
V = lambda sh, a: wv[sh][a].value
NIPPO = "日報（認定）"
days = [V(NIPPO, f"A{50*k+1}") for k in range(31)]
days = [(k, d) for k, d in enumerate(days) if isinstance(d, datetime.datetime)]
MONTH = days[0][1].strftime("%Y%m")
YM = days[0][1].strftime("%Y/%m")
LABEL = f"{days[0][1].year}年{days[0][1].month}月"

FONT = "Meiryo UI"; NAVY = "16324F"
F_HEAD = PatternFill("solid", fgColor=NAVY)
THIN = Side(style="thin", color="C9D6DD"); BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
def font(**k): return Font(name=FONT, **{"size": 10, **k})
def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = font(size=16, bold=True, color=NAVY)
    if sub: ws["A2"] = sub; ws["A2"].font = font(size=10, color="5B6B77")
    ws.sheet_view.showGridLines = False
def table(ws, r0, headers, rows, widths, wrap_cols=()):
    for j, h in enumerate(headers, 1):
        c = ws.cell(r0, j, h); c.font = font(bold=True, color="FFFFFF"); c.fill = F_HEAD; c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(r0 + i, j, v); c.font = font(); c.border = BOX
            c.alignment = Alignment(vertical="top", wrap_text=(j in wrap_cols))
    for j, w in enumerate(widths, 1): ws.column_dimensions[gl(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1)
def wcsv(name, headers, rows):
    with open(os.path.join(CSVD, name), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(headers); w.writerows(rows)

# ---------------- マスタ ----------------
M = {}
M["M01_直"] = (["直コード", "直名称", "開始時刻", "終了時刻", "標準所定時間分", "使用中"],
              [["D", "D直", "07:00", "19:00", 720, "○"], ["E", "E直", "15:00", "23:00", 480, "×（現在未使用）"],
               ["N", "N直", "19:00", "07:00", 720, "○"]])
M["M02_班"] = (["班コード", "班名称"], [["A", "A班"], ["B", "B班"], ["C", "C班"]])
M["M03_搬入元"] = (["搬入元コード", "搬入元名", "原料区分", "使用中"],
                  [["S01", "仙台市（大）", "認定", "○"], ["S02", "仙台市（小）", "認定", "○"], ["S03", "札幌市", "容リ", "○"],
                   ["S04", "独自ベール（川崎・太田区）", "独自", "○"], ["S05", "東部衛生処理組合", "32条", "○"]])
M["M04_品目"] = (["品目コード", "品目名", "区分", "標準単体重量kg", "出荷シートでの名称", "備考"],
                [["P01", "PE･PP混合減容品（フレコン）", "製品", "", "PE･PP混合減容品・フラフ", "重量はフレコン明細の合計"],
                 ["P02", "PE･PP混合減容品（直送）", "製品", "", "PEPP混合減容品（直送）", "直送ライン計量"],
                 ["P03", "PE･PP混合フラフ", "製品", "", "PE･PP混合減容品・フラフ", "直送ライン計量"],
                 ["P04", "PE・PP混合フレーク", "製品", "", "PE・PP混合フレーク", ""],
                 ["P05", "PEフレーク", "製品", "", "PEフレーク", ""],
                 ["P06", "PPフレーク", "製品", "", "PPフレーク", ""],
                 ["P07", "PSフレーク", "製品", "", "PSフレーク", ""],
                 ["P08", "PETフレーク", "製品", "", "PETフレーク", ""],
                 ["P09", "PSインゴット", "製品", "", "PSインゴット", ""],
                 ["W01", "残渣ベール", "廃棄物", 185, "残渣ベール", "標準単体重量は旧日報T列（前月ファイル参照）"],
                 ["W02", "洗浄残渣（フレコン）", "廃棄物", 280, "洗浄残渣", ""],
                 ["W03", "残渣フレコン（Z）", "廃棄物", 55, "残渣フレコン", "旧日報では前月値-5kg"],
                 ["W04", "不適物", "廃棄物", "", "手選別不適物", "重量は実測で入力"],
                 ["W05", "脱水汚泥", "廃棄物", "", "脱水汚泥", "在庫は日報で入力"]])
dest, grp = [], None
for c in range(2, 159):
    L = gl(c); h3, h4 = V("出荷", f"{L}3"), V("出荷", f"{L}4")
    if h3: grp = str(h3)
    if h4 in (None, 0, "0", "#REF!", "合計", "外販小計") or grp == "PE･PP": continue
    dest.append([grp, str(h4)])
M["M05_出荷先"] = (["出荷先コード", "出荷先名", "対象品目（出荷シートの区分）", "備考"],
                  [[f"C{i:02d}", n, g, "パレット出荷は他ファイル（原料）から連携" if "パレット" in n else ""] for i, (g, n) in enumerate(dest, 1)])
eq = [str(V(NIPPO, f"X{r}")).lstrip("・") for r in range(30, 39)]
M["M06_設備"] = (["設備コード", "設備名", "実操業時間の判定に使う", "表示順"],
                [[f"E{i:02d}", n, "○" if n in ("開梱機", "破袋機") else "", i] for i, n in enumerate(eq, 1)])
pr = [str(V(NIPPO, f"X{r}")) for r in range(18, 30)]
M["M07_工程"] = (["工程コード", "工程名", "表示順"], [[f"K{i:02d}", n, i] for i, n in enumerate(pr, 1)])
M["M08_休止要因"] = (["要因コード", "要因名", "説明"],
                    [["R1", "人的", "欠員・休憩ずらし等"], ["R2", "原料", "巻き付き・異物・ベール性状"], ["R3", "設備", "故障・詰まり・調整"], ["R4", "その他", "訓練・立上げ/立下げ等"]])
cal = []
for k, dt in days:
    y, a = V(NIPPO, f"Y{50*k+8}") or 0, V(NIPPO, f"AA{50*k+8}") or 0
    t = V(NIPPO, f"Y{50*k+16}") or 0
    kind = "休止" if y == 0 and a == 0 else ("定修あり" if t else "通常")
    cal.append([dt.strftime("%Y/%m/%d"), "月火水木金土日"[dt.weekday()], kind, y, a])
for name, (h, rows) in M.items():
    wcsv(f"{name}.csv", h, rows)
# M09は月ごとに増えるのでファイル名に年月を付ける
wcsv(f"M09_操業カレンダー_{MONTH}.csv", ["日付", "曜日", "区分", "D直所定時間分", "N直所定時間分"], cal)

# ---------------- T90 月初在庫（旧：前月ファイル参照） ----------------
inv = []
for sh, rows in ((NIPPO, list(range(9, 12)) + list(range(22, 30)) + list(range(32, 37))), ("日報（容リ）", range(12, 15)), ("日報（32条）", range(15, 17))):
    for r in rows:
        name = V(sh, f"C{r}")
        if name in (None, 0, "0"): continue
        w, n = V(sh, f"E{r}"), V(sh, f"F{r}")
        kubun = "原料" if r < 20 else ("製品" if r < 31 else "廃棄物")
        inv.append([YM, kubun, name, w or 0, n if isinstance(n, (int, float)) else ""])
wcsv(f"T90_月初在庫_{MONTH}.csv", ["年月", "区分", "品目・搬入元", "重量kg", "個数"], inv)

# ---------------- 検証（旧日報の集計値 vs 移行CSVの再集計） ----------------
xsrc = os.path.join(CSVD, f"X_検証_日報集計値_{MONTH}.csv")
xdst = os.path.join(CSVD, f"99_検証_旧日報の日別集計値_{MONTH}.csv")
os.replace(xsrc, xdst)
Lr = lambda p: pd.read_csv(os.path.join(CSVD, f"{p}_{MONTH}.csv"))
x = pd.read_csv(xdst)
waste = {"残渣ベール", "洗浄残渣", "残渣フレコン", "手選別不適物", "脱水汚泥"}
s7 = Lr("T07_出荷明細"); s7["w"] = s7["品目"].isin(waste)
t9 = Lr("T09_設備稼働"); t9["直"] = t9["時間帯"].str[0]
op = t9[t9["設備"].isin(["開梱機", "破袋機"])].groupby(["日付", "直", "設備"])["稼働分"].sum().unstack().max(axis=1).groupby("日付").sum()
checks = [
 ("原料入荷 重量kg", x["原料入荷kg"].sum(), Lr("T04_入荷明細")["重量kg"].sum(), "T04_入荷明細"),
 ("原料入荷 個数", x["原料入荷個数"].sum(), Lr("T04_入荷明細")["個数"].sum(), "T04_入荷明細"),
 ("原料処理 個数", x["原料処理個数"].sum(), Lr("T03_原料処理")["処理個数"].sum(), "T03_原料処理"),
 ("原料処理 重量kg", x["原料処理kg"].sum(), Lr("T03_原料処理")["処理重量kg"].sum(), "T03_原料処理"),
 ("減容品 生産重量kg", x["減容品生産kg"].sum(), Lr("T05_減容品フレコン")["重量kg"].sum(), "T05_減容品フレコン"),
 ("減容品 袋数", x["減容品生産袋数"].sum(), len(Lr("T05_減容品フレコン")), "T05_減容品フレコン"),
 ("製品 出荷kg", x["製品出荷kg"].sum(), s7[~s7.w]["重量kg"].sum(), "T07_出荷明細"),
 ("廃棄物 出荷kg", x["廃棄物出荷kg"].sum(), s7[s7.w]["重量kg"].sum(), "T07_出荷明細"),
 ("実操業時間 分", x["実操業時間分"].sum(), op.sum(), "T09_設備稼働"),
]
ng = [(a, b, c) for a, b, c, _ in checks if abs(float(b) - float(c)) > 0.5]

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "移行手順"
title(ws, "過去データ移行の手順（CSV取込）")
STEPS = [
 ("1", "マスタを先に取り込む", "M01〜M09 のCSVを取り込む。参照項目はマスタが先にないとエラーになる"),
 ("2", "T90_月初在庫を取り込む", "在庫計算の起点。前月ファイルへのリンクの代わり"),
 ("3", "T01_日報 → T02_直別操業 の順に取り込む", "他の記録テーブルが参照する親テーブル"),
 ("4", "残りのT03〜T12を取り込む", "順番は自由"),
 ("5", "件数と合計を確認する", "「検証結果」シートの数字（旧日報の集計値）とJUST.DBの集計が一致するか確認"),
]
table(ws, 4, ["手順", "作業", "ポイント"], STEPS, [6, 36, 80], wrap_cols=(3,))
ws.cell(11, 1, "CSVの形式：UTF-8（BOM付き）、1行目が項目名、日付は yyyy/mm/dd、時刻は hh:mm。").font = font(color="5B6B77")

ws = wb.create_sheet("ファイル一覧")
title(ws, "移行用CSV のファイル一覧")
files = []
for p in sorted(glob.glob(os.path.join(CSVD, "*.csv"))):
    with open(p, encoding="utf-8-sig") as fh: files.append((os.path.basename(p), sum(1 for _ in fh) - 1))
table(ws, 4, ["ファイル名", "件数"], files, [48, 10])

ws = wb.create_sheet("検証結果")
title(ws, f"検証結果：移行CSVから再集計した値と旧日報の値の比較（{LABEL}）",
      "すべての項目で差は0でした。" if not ng else "差がある項目があります。差の列を確認してください。")
rows_ = [(a, float(b), float(c), f"=C{5+i}-B{5+i}", t) for i, (a, b, c, t) in enumerate(checks)]
table(ws, 4, [f"項目（{LABEL}合計）", "旧日報の値", "移行CSVから再集計", "差", "元のCSV"], rows_, [22, 16, 18, 10, 22])
for r in range(5, 5 + len(rows_)):
    for c in (2, 3, 4): ws.cell(r, c).number_format = "#,##0"
wb.save(os.path.join(CSVD, f"00_移行手順と検証結果_{MONTH}.xlsx"))

for a, b, c, _ in checks:
    print(f"{'OK' if abs(float(b)-float(c)) <= 0.5 else 'NG'}  {a}: 旧日報 {float(b):,.0f} / CSV {float(c):,.0f}")
sys.exit(1 if ng else 0)
