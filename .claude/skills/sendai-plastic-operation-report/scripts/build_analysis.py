import os, csv, glob, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as gl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.workbook.defined_name import DefinedName

"""JUST.DBのCSV（T02〜T12）から、CSV差し替え式の分析Excelを作る
使い方: python build_analysis.py <CSVフォルダ> <出力.xlsx>
CSVフォルダに同じテーブルのCSVが複数月ある場合は、すべて結合して取り込む。
"""
import sys
CSVD, OUT = sys.argv[1], sys.argv[2]
FONT = "Meiryo UI"
NAVY, S1, S2, MUTED, GRID = "16324F", "0B7FA0", "D0661C", "5B6B77", "D9E1E6"
F_HEAD = PatternFill("solid", fgColor=NAVY); F_IN = PatternFill("solid", fgColor="FFF2CC")
F_TOT = PatternFill("solid", fgColor="E8EEF2"); F_TILE = PatternFill("solid", fgColor="F1F6F8")
THIN = Side(style="thin", color="C9D6DD"); BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
def font(**k): return Font(name=FONT, **{"size": 10, **k})

wb = openpyxl.Workbook()
use = wb.active; use.title = "使い方"
dash = wb.create_sheet("ダッシュボード")
day = wb.create_sheet("日別集計")
shf = wb.create_sheet("直別集計")
stop = wb.create_sheet("休止・トラブル")
ship = wb.create_sheet("出荷集計")

# ---------------- データシート（CSVをそのまま貼る） ----------------
DATA = ["T02_直別操業", "T03_原料処理", "T04_入荷明細", "T05_減容品フレコン", "T06_廃棄物発生",
        "T07_出荷明細", "T08_予定外休止", "T09_設備稼働", "T10_特記事項", "T12_直送時間別"]
def conv(k, v):
    if v == "": return None
    if k == "日付":
        try: return datetime.datetime.strptime(v, "%Y/%m/%d")
        except ValueError: return v
    try: return int(v)
    except ValueError:
        try: return float(v)
        except ValueError: return v
for t in DATA:
    ws = wb.create_sheet(t)
    ws.sheet_properties.tabColor = "A6B4BE"
    hdr = None
    for p in sorted(glob.glob(os.path.join(CSVD, f"{t}_*.csv"))):
        with open(p, encoding="utf-8-sig") as f:
            rd = csv.reader(f); h = next(rd)
            if hdr is None: hdr = h; ws.append(hdr)
            for row in rd:
                rec = dict(zip(h, row))
                ws.append([conv(k, rec.get(k, "")) for k in hdr])
    if hdr is None:
        sys.exit(f"{t} のCSVがありません。extract.py で変換し直してください（0件の月も見出しだけのCSVが出ます）")
    for j in range(1, len(hdr) + 1):
        ws.cell(1, j).font = font(bold=True, color="FFFFFF"); ws.cell(1, j).fill = F_HEAD
        ws.column_dimensions[gl(j)].width = 14
    for r in range(2, ws.max_row + 1): ws.cell(r, 1).number_format = "yyyy/mm/dd"
    ws.freeze_panes = "A2"

def C(t, name):
    """見出し名で列を探す（CSVの列の順番が変わっても動く）"""
    return f'INDEX({t}!$A$1:$Z$30000,0,MATCH("{name}",{t}!$A$1:$Z$1,0))'
MON = '">="&開始日,{d},"<"&終了日'
def inmonth(t): return f'{C(t,"日付")},">="&開始日,{C(t,"日付")},"<"&終了日'

# ---------------- 使い方 ----------------
use.sheet_view.showGridLines = False
use["A1"] = "操業分析（CSV差し替え式）"; use["A1"].font = font(size=18, bold=True, color=NAVY)
use["A3"] = "対象月（1日の日付を入力）"; use["A3"].font = font(bold=True)
_last = max(c.value for n in DATA for c in wb[n]["A"][1:] if isinstance(c.value, datetime.datetime))
use["C3"] = datetime.datetime(_last.year, _last.month, 1); use["C3"].number_format = "yyyy年m月"; use["C3"].fill = F_IN
use["C3"].font = font(size=12, bold=True); use["C3"].border = BOX
use["D3"] = "← 黄色のセルだけ変更します"; use["D3"].font = font(color=MUTED)
wb.defined_names["開始日"] = DefinedName("開始日", attr_text="使い方!$C$3")
wb.defined_names["終了日"] = DefinedName("終了日", attr_text="使い方!$C$4")
use["A4"] = "（翌月1日・自動）"; use["A4"].font = font(color=MUTED)
use["C4"] = "=EDATE(C3,1)"; use["C4"].number_format = "yyyy/mm/dd"; use["C4"].font = font(color=MUTED)
steps = [
 "【更新のしかた】",
 "1. JUST.DBで各テーブルをCSV出力する（テーブル名はシート名と同じ：T02_直別操業 など灰色タブの10シート）。",
 "2. CSVをExcelで開き、全体をコピーする。",
 "3. このファイルの同じ名前のシートを開き、A1セルを選んで貼り付ける（前のデータは先に全部消す）。",
 "4. 上の「対象月」を変える。集計・グラフは自動で更新されます。",
 "",
 "【ポイント】",
 "・1行目の見出し名で列を探しているので、列の順番が変わっても動きます。見出し名は変えないでください。",
 "・1年分など複数月のデータを貼っても、対象月の分だけ集計します（各シート30,000行まで）。",
 "・原料処理重量（処理重量kg）と廃棄物の発生重量（発生重量kg）はJUST.DB側の計算項目の値を使います。",
 "・実操業時間は「開梱機と破袋機の稼働時間の長いほう」（旧日報と同じルール）で計算しています。",
 "・製品生産量＝減容品（フレコン）＋減容品（直送）＋フラフ。製品収率＝製品生産量÷原料処理重量。",
 "",
 "【シート】",
 "ダッシュボード：月の主要指標とグラフ ／ 日別集計：1日1行の集計表 ／ 直別集計：D直とN直の比較",
 "休止・トラブル：工程別の予定外休止、要因別のトラブル、設備別の稼働 ／ 出荷集計：品目・出荷先別",
 "灰色タブ：CSVを貼るデータシート",
]
for i, s in enumerate(steps, 6):
    use.cell(i, 1, s).font = font(bold=s.startswith("【"), color=NAVY if s.startswith("【") else "1E2A33")
use.column_dimensions["A"].width = 26; use.column_dimensions["C"].width = 14

# ---------------- 日別集計 ----------------
day.sheet_view.showGridLines = False
day["A1"] = "日別集計"; day["A1"].font = font(size=16, bold=True, color=NAVY)
day["A2"] = '=TEXT(開始日,"yyyy年m月")&"　（kg・分）"'; day["A2"].font = font(color=MUTED)
WASTE = '{"残渣ベール","洗浄残渣","残渣フレコン","手選別不適物","脱水汚泥"}'
cols = [
 ("日付", None, "m/d(aaa)"),
 ("原料入荷kg", lambda r: f'SUMIFS({C("T04_入荷明細","重量kg")},{C("T04_入荷明細","日付")},$A{r})', "#,##0"),
 ("処理個数", lambda r: f'SUMIFS({C("T03_原料処理","処理個数")},{C("T03_原料処理","日付")},$A{r})', "#,##0"),
 ("原料処理kg", lambda r: f'SUMIFS({C("T03_原料処理","処理重量kg")},{C("T03_原料処理","日付")},$A{r})', "#,##0"),
 ("減容品kg", lambda r: f'SUMIFS({C("T05_減容品フレコン","重量kg")},{C("T05_減容品フレコン","日付")},$A{r})', "#,##0"),
 ("減容品袋数", lambda r: f'COUNTIFS({C("T05_減容品フレコン","日付")},$A{r})', "#,##0"),
 ("直送減容品kg", lambda r: f'ROUND(SUMIFS({C("T12_直送時間別","重量kg")},{C("T12_直送時間別","日付")},$A{r},{C("T12_直送時間別","品目")},"PE･PP混合減容品（直送）"),0)', "#,##0"),
 ("フラフkg", lambda r: f'ROUND(SUMIFS({C("T12_直送時間別","重量kg")},{C("T12_直送時間別","日付")},$A{r},{C("T12_直送時間別","品目")},"PE･PP混合フラフ"),0)', "#,##0"),
 ("製品生産kg", lambda r: f"E{r}+G{r}+H{r}", "#,##0"),
 ("製品収率", lambda r: f'IF(D{r}=0,"",I{r}/D{r})', "0.0%"),
 ("廃棄物発生kg", lambda r: f'SUMIFS({C("T06_廃棄物発生","発生重量kg")},{C("T06_廃棄物発生","日付")},$A{r})', "#,##0"),
 ("廃棄物比率", lambda r: f'IF(D{r}=0,"",K{r}/D{r})', "0.0%"),
 ("実操業D分", lambda r: f'MAX(SUMIFS({C("T09_設備稼働","稼働分")},{C("T09_設備稼働","日付")},$A{r},{C("T09_設備稼働","時間帯")},"D",{C("T09_設備稼働","設備")},"開梱機"),SUMIFS({C("T09_設備稼働","稼働分")},{C("T09_設備稼働","日付")},$A{r},{C("T09_設備稼働","時間帯")},"D",{C("T09_設備稼働","設備")},"破袋機"))', "#,##0"),
 ("実操業N分", lambda r: f'MAX(SUMIFS({C("T09_設備稼働","稼働分")},{C("T09_設備稼働","日付")},$A{r},{C("T09_設備稼働","時間帯")},"N*",{C("T09_設備稼働","設備")},"開梱機"),SUMIFS({C("T09_設備稼働","稼働分")},{C("T09_設備稼働","日付")},$A{r},{C("T09_設備稼働","時間帯")},"N*",{C("T09_設備稼働","設備")},"破袋機"))', "#,##0"),
 ("実操業計分", lambda r: f"M{r}+N{r}", "#,##0"),
 ("処理ペースkg/h", lambda r: f'IF(O{r}=0,"",D{r}/O{r}*60)', "#,##0"),
 ("予定外休止分", lambda r: f'SUMIFS({C("T08_予定外休止","休止分")},{C("T08_予定外休止","日付")},$A{r})', "#,##0"),
 ("減容品平均kg/袋", lambda r: f'IF(F{r}=0,"",E{r}/F{r})', "#,##0"),
 ("出荷計kg", lambda r: f'SUMIFS({C("T07_出荷明細","重量kg")},{C("T07_出荷明細","日付")},$A{r})', "#,##0"),
 ("うち廃棄物kg", lambda r: f'SUMPRODUCT(SUMIFS({C("T07_出荷明細","重量kg")},{C("T07_出荷明細","日付")},$A{r},{C("T07_出荷明細","品目")},{WASTE}))', "#,##0"),
 ("トラブル件数", lambda r: f'COUNTIFS({C("T10_特記事項","日付")},$A{r},{C("T10_特記事項","要因")},"?*")', "#,##0"),
 ("トラブル分", lambda r: f'SUMIFS({C("T10_特記事項","時間分")},{C("T10_特記事項","日付")},$A{r})', "#,##0"),
]
R0 = 4
for j, (h, fn, fmt) in enumerate(cols, 1):
    c = day.cell(R0, j, h); c.font = font(bold=True, color="FFFFFF"); c.fill = F_HEAD
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOX
    day.column_dimensions[gl(j)].width = 11 if j > 1 else 10
day.row_dimensions[R0].height = 30
FIRST, LAST = R0 + 1, R0 + 31
for i in range(31):
    r = FIRST + i
    day.cell(r, 1, f'=IF(開始日+{i}<終了日,開始日+{i},"")')
    for j, (h, fn, fmt) in enumerate(cols, 1):
        c = day.cell(r, j)
        if fn: c.value = f'=IF($A{r}="","",{fn(r)})'
        c.number_format = fmt; c.font = font(); c.border = BOX
TOT = LAST + 1
day.cell(TOT, 1, "月計").font = font(bold=True)
avg_cols = {"製品収率": f"IF(D{TOT}=0,\"\",I{TOT}/D{TOT})", "廃棄物比率": f"IF(D{TOT}=0,\"\",K{TOT}/D{TOT})",
            "処理ペースkg/h": f"IF(O{TOT}=0,\"\",D{TOT}/O{TOT}*60)", "減容品平均kg/袋": f"IF(F{TOT}=0,\"\",E{TOT}/F{TOT})"}
for j, (h, fn, fmt) in enumerate(cols, 1):
    c = day.cell(TOT, j); c.fill = F_TOT; c.border = BOX; c.font = font(bold=True); c.number_format = fmt
    if j > 1: c.value = "=" + avg_cols[h] if h in avg_cols else f"=SUM({gl(j)}{FIRST}:{gl(j)}{LAST})"
day.freeze_panes = day.cell(FIRST, 2)
COL = {h: gl(j) for j, (h, _, _) in enumerate(cols, 1)}

# ---------------- 直別集計 ----------------
shf.sheet_view.showGridLines = False
shf["A1"] = "直別集計（D直・N直の比較）"; shf["A1"].font = font(size=16, bold=True, color=NAVY)
shf["A2"] = '=TEXT(開始日,"yyyy年m月")'; shf["A2"].font = font(color=MUTED)
def mshift(t, field, sh, how="SUMIFS"):
    return f'{how}({C(t,field)},{inmonth(t)},{C(t,"直")},"{sh}")'
rows = [
 ("原料処理kg", lambda s: mshift("T03_原料処理", "処理重量kg", s), "#,##0"),
 ("処理個数", lambda s: mshift("T03_原料処理", "処理個数", s), "#,##0"),
 ("減容品kg", lambda s: mshift("T05_減容品フレコン", "重量kg", s), "#,##0"),
 ("減容品袋数", lambda s: f'COUNTIFS({inmonth("T05_減容品フレコン")},{C("T05_減容品フレコン","直")},"{s}")', "#,##0"),
 ("減容品収率（減容品÷処理）", lambda s: None, "0.0%"),
 ("廃棄物発生kg", lambda s: mshift("T06_廃棄物発生", "発生重量kg", s), "#,##0"),
 ("実操業時間分", lambda s: f"日別集計!{COL['実操業D分' if s=='D' else '実操業N分']}{TOT}", "#,##0"),
 ("処理ペースkg/h", lambda s: None, "#,##0"),
 ("予定外休止分", lambda s: mshift("T08_予定外休止", "休止分", s), "#,##0"),
 ("トラブル件数", lambda s: f'COUNTIFS({inmonth("T10_特記事項")},{C("T10_特記事項","記入欄")},"{s}直側",{C("T10_特記事項","要因")},"?*")', "#,##0"),
 ("操業した直の数", lambda s: f'COUNTIFS({inmonth("T02_直別操業")},{C("T02_直別操業","直")},"{s}")', "#,##0"),
]
for j, h in enumerate(["指標", "D直", "N直", "合計"], 1):
    c = shf.cell(4, j, h); c.font = font(bold=True, color="FFFFFF"); c.fill = F_HEAD; c.border = BOX
    c.alignment = Alignment(horizontal="center")
for i, (name, fn, fmt) in enumerate(rows, 5):
    shf.cell(i, 1, name).font = font(bold=True)
    for j, s in ((2, "D"), (3, "N")):
        f = fn(s)
        if name.startswith("減容品収率"): f = f'IF({gl(j)}5=0,"",{gl(j)}7/{gl(j)}5)'
        if name.startswith("処理ペース"): f = f'IF({gl(j)}11=0,"",{gl(j)}5/{gl(j)}11*60)'
        shf.cell(i, j, "=" + f)
    shf.cell(i, 4, f'=IF(B{i}="","",B{i}+C{i})' if "収率" not in name and "ペース" not in name else
             (f'IF(D5=0,"",D7/D5)' if "収率" in name else f'IF(D11=0,"",D5/D11*60)'))
    if not str(shf.cell(i, 4).value).startswith("="): shf.cell(i, 4).value = "=" + shf.cell(i, 4).value
    for j in range(1, 5):
        c = shf.cell(i, j); c.border = BOX; c.number_format = fmt
        if j > 1: c.font = font()
shf.column_dimensions["A"].width = 28
for L in "BCD": shf.column_dimensions[L].width = 14

# ---------------- 休止・トラブル ----------------
stop.sheet_view.showGridLines = False
stop["A1"] = "休止・トラブル・設備稼働"; stop["A1"].font = font(size=16, bold=True, color=NAVY)
stop["A2"] = '=TEXT(開始日,"yyyy年m月")&"　（分）"'; stop["A2"].font = font(color=MUTED)
def read_master(n):
    with open(os.path.join(CSVD, f"{n}.csv"), encoding="utf-8-sig") as f: return list(csv.DictReader(f))
def block(ws, r0, c0, head, items, fns, fmts):
    for j, h in enumerate(head):
        c = ws.cell(r0, c0 + j, h); c.font = font(bold=True, color="FFFFFF"); c.fill = F_HEAD; c.border = BOX
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, it in enumerate(items, 1):
        ws.cell(r0 + i, c0, it).border = BOX; ws.cell(r0 + i, c0).font = font()
        for j, fn in enumerate(fns, 1):
            c = ws.cell(r0 + i, c0 + j, "=" + fn(r0 + i, it)); c.border = BOX; c.font = font(); c.number_format = fmts[j - 1]
    return r0 + len(items)
procs = [m["工程名"] for m in read_master("M07_工程") if m["工程名"] != "その他"]
stop.cell(4, 1, "工程別 予定外休止（分）").font = font(size=12, bold=True, color=NAVY)
pe = block(stop, 5, 1, ["工程", "D直", "N直", "合計"], procs,
           [lambda r, it: f'SUMIFS({C("T08_予定外休止","休止分")},{inmonth("T08_予定外休止")},{C("T08_予定外休止","工程")},$A{r},{C("T08_予定外休止","直")},"D")',
            lambda r, it: f'SUMIFS({C("T08_予定外休止","休止分")},{inmonth("T08_予定外休止")},{C("T08_予定外休止","工程")},$A{r},{C("T08_予定外休止","直")},"N")',
            lambda r, it: f"B{r}+C{r}"], ["#,##0"] * 3)
facs = [m["要因名"] for m in read_master("M08_休止要因")]
stop.cell(4, 6, "要因別 トラブル・特記事項").font = font(size=12, bold=True, color=NAVY)
fe = block(stop, 5, 6, ["要因", "件数", "時間（分）"], facs,
           [lambda r, it: f'COUNTIFS({inmonth("T10_特記事項")},{C("T10_特記事項","要因")},$F{r})',
            lambda r, it: f'SUMIFS({C("T10_特記事項","時間分")},{inmonth("T10_特記事項")},{C("T10_特記事項","要因")},$F{r})'], ["#,##0"] * 2)
eqs = [m["設備名"] for m in read_master("M06_設備")]
stop.cell(fe + 3, 6, "設備別 稼働時間（時間）").font = font(size=12, bold=True, color=NAVY)
block(stop, fe + 4, 6, ["設備", "D直", "N直", "合計"], eqs,
      [lambda r, it: f'SUMIFS({C("T09_設備稼働","稼働分")},{inmonth("T09_設備稼働")},{C("T09_設備稼働","設備")},$F{r},{C("T09_設備稼働","時間帯")},"D")/60',
       lambda r, it: f'SUMIFS({C("T09_設備稼働","稼働分")},{inmonth("T09_設備稼働")},{C("T09_設備稼働","設備")},$F{r},{C("T09_設備稼働","時間帯")},"N*")/60',
       lambda r, it: f"G{r}+H{r}"], ["#,##0.0"] * 3)
for L, w in zip("ABCDEFGHI", [22, 9, 9, 9, 3, 22, 9, 11, 9]): stop.column_dimensions[L].width = w

# ---------------- 出荷集計 ----------------
ship.sheet_view.showGridLines = False
ship["A1"] = "出荷集計（品目・出荷先別）"; ship["A1"].font = font(size=16, bold=True, color=NAVY)
ship["A2"] = '=TEXT(開始日,"yyyy年m月")'; ship["A2"].font = font(color=MUTED)
dests = [(m["対象品目（出荷シートの区分）"], m["出荷先名"]) for m in read_master("M05_出荷先")]
for j, h in enumerate(["品目", "出荷先", "重量kg", "個数", "出荷回数"], 1):
    c = ship.cell(4, j, h); c.font = font(bold=True, color="FFFFFF"); c.fill = F_HEAD; c.border = BOX
for i, (g, n) in enumerate(dests, 5):
    ship.cell(i, 1, g); ship.cell(i, 2, n)
    crit = f'{inmonth("T07_出荷明細")},{C("T07_出荷明細","品目")},$A{i},{C("T07_出荷明細","出荷先")},$B{i}'
    ship.cell(i, 3, f'=SUMIFS({C("T07_出荷明細","重量kg")},{crit})')
    ship.cell(i, 4, f'=SUMIFS({C("T07_出荷明細","個数")},{crit})')
    ship.cell(i, 5, f'=COUNTIFS({crit})')
    for j in range(1, 6):
        c = ship.cell(i, j); c.border = BOX; c.font = font(); c.number_format = "#,##0"
SE = 4 + len(dests)
ship.cell(SE + 1, 2, "合計").font = font(bold=True)
for j in (3, 4, 5):
    c = ship.cell(SE + 1, j, f"=SUM({gl(j)}5:{gl(j)}{SE})"); c.font = font(bold=True); c.fill = F_TOT; c.number_format = "#,##0"
ship.auto_filter.ref = f"A4:E{SE}"
for L, w in zip("ABCDE", [26, 26, 12, 10, 10]): ship.column_dimensions[L].width = w
ship.freeze_panes = "A5"

# ---------------- ダッシュボード ----------------
dash.sheet_view.showGridLines = False
dash["B1"] = '="操業ダッシュボード　"&TEXT(開始日,"yyyy年m月")'; dash["B1"].font = font(size=18, bold=True, color=NAVY)
dash["B2"] = "対象月は「使い方」シートで変更"; dash["B2"].font = font(color=MUTED)
T = f"日別集計!{{}}{TOT}"
tiles = [
 ("原料処理量", f"=日別集計!{COL['原料処理kg']}{TOT}/1000", '#,##0.0" t"'),
 ("製品生産量", f"=日別集計!{COL['製品生産kg']}{TOT}/1000", '#,##0.0" t"'),
 ("製品収率", f"=日別集計!{COL['製品収率']}{TOT}", "0.0%"),
 ("廃棄物比率", f"=日別集計!{COL['廃棄物比率']}{TOT}", "0.0%"),
 ("実操業時間", f"=日別集計!{COL['実操業計分']}{TOT}/60", '#,##0" h"'),
 ("処理ペース", f"=日別集計!{COL['処理ペースkg/h']}{TOT}", '#,##0" kg/h"'),
 ("予定外休止", f"=日別集計!{COL['予定外休止分']}{TOT}/60", '#,##0.0" h"'),
 ("トラブル件数", f"=日別集計!{COL['トラブル件数']}{TOT}", '#,##0" 件"'),
]
for i, (lab, f, fmt) in enumerate(tiles):
    col = 2 + (i % 4) * 3; row = 4 + (i // 4) * 3
    for rr in (row, row + 1):
        for cc in (col, col + 1):
            dash.cell(rr, cc).fill = F_TILE
    dash.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
    dash.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
    c = dash.cell(row, col, lab); c.font = font(color=MUTED); c.alignment = Alignment(horizontal="left", indent=1)
    c = dash.cell(row + 1, col, f); c.font = font(size=18, bold=True, color=NAVY); c.number_format = fmt
    c.alignment = Alignment(horizontal="left", indent=1)
    dash.row_dimensions[row + 1].height = 30
for L in "BCDEFGHIJKLM": dash.column_dimensions[L].width = 11
for L in "DGJ": dash.column_dimensions[L].width = 3

def style_axes(ch, ytitle):
    ch.y_axis.title = ytitle
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=GRID))
    ch.y_axis.delete = False; ch.x_axis.delete = False
    ch.x_axis.number_format = "m/d"
    ch.height, ch.width = 7.5, 24
cats = Reference(day, min_col=1, min_row=FIRST, max_row=LAST)
# 1) 日別 原料処理量と製品生産量（2系列・棒）
b = BarChart(); b.type = "col"; b.grouping = "clustered"; b.gapWidth = 60; b.overlap = -10
b.title = "日別 原料処理量と製品生産量（kg）"
for colname, color in (("原料処理kg", S1), ("製品生産kg", S2)):
    j = list(COL).index(colname) + 1
    b.add_data(Reference(day, min_col=j, min_row=R0, max_row=LAST), titles_from_data=True)
    s = b.series[-1]; s.graphicalProperties.solidFill = color; s.graphicalProperties.line.noFill = True
b.set_categories(cats); style_axes(b, "kg"); b.legend.position = "t"
dash.add_chart(b, "B11")
# 2) 日別 製品収率（1系列・棒、凡例なし。非稼働日は棒なし）
ln = BarChart(); ln.type = "col"; ln.gapWidth = 60; ln.title = "日別 製品収率"
j = list(COL).index("製品収率") + 1
ln.add_data(Reference(day, min_col=j, min_row=R0, max_row=LAST), titles_from_data=True)
s = ln.series[0]; s.graphicalProperties.solidFill = S1; s.graphicalProperties.line.noFill = True
ln.set_categories(cats); style_axes(ln, "収率"); ln.y_axis.number_format = "0%"; ln.legend = None
dash.add_chart(ln, "B27")
# 3) 工程別 予定外休止（1系列・横棒）
hb = BarChart(); hb.type = "bar"; hb.title = "工程別 予定外休止（分）"; hb.gapWidth = 50
hb.add_data(Reference(stop, min_col=4, min_row=5, max_row=pe), titles_from_data=True)
hb.set_categories(Reference(stop, min_col=1, min_row=6, max_row=pe))
hb.series[0].graphicalProperties.solidFill = S1; hb.series[0].graphicalProperties.line.noFill = True
hb.legend = None; hb.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=GRID))
hb.x_axis.scaling.orientation = "maxMin"; hb.y_axis.delete = False; hb.x_axis.delete = False
hb.height, hb.width = 8.5, 11.5
dash.add_chart(hb, "B43")
# 4) 要因別 トラブル件数（1系列・横棒）
fb = BarChart(); fb.type = "bar"; fb.title = "要因別 トラブル・特記事項（件）"; fb.gapWidth = 50
fb.add_data(Reference(stop, min_col=7, min_row=5, max_row=fe), titles_from_data=True)
fb.set_categories(Reference(stop, min_col=6, min_row=6, max_row=fe))
fb.series[0].graphicalProperties.solidFill = S1; fb.series[0].graphicalProperties.line.noFill = True
fb.legend = None; fb.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=GRID))
fb.x_axis.scaling.orientation = "maxMin"; fb.y_axis.delete = False; fb.x_axis.delete = False
fb.height, fb.width = 8.5, 11.5
dash.add_chart(fb, "H43")

wb.move_sheet("使い方", offset=0)
wb.active = 1
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved")
