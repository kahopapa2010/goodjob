"""操業日報Excel（月次ファイル）→ JUST.DB移行用CSV 変換スクリプト
使い方: python extract.py <月次日報.xlsx> <出力フォルダ>
"""
import sys, os, csv, datetime, warnings, re
import openpyxl
from openpyxl.utils import get_column_letter as gl, column_index_from_string as ci
warnings.filterwarnings("ignore")

src, outdir = sys.argv[1], sys.argv[2]
os.makedirs(outdir, exist_ok=True)
wf = openpyxl.load_workbook(src)                  # 数式
wv = openpyxl.load_workbook(src, data_only=True)  # 計算結果

CONST = re.compile(r"^=[\d.+\-*/() ]+$")   # 例: =18+24（セル内で手計算した入力）
def isf(ws, a):
    v = wf[ws][a].value
    return isinstance(v, str) and v.startswith("=") and not CONST.match(v)
def const_parts(ws, a):
    """=12290+11010 のような足し算入力を個々の値に分解（それ以外は None）"""
    v = wf[ws][a].value
    if isinstance(v, str) and re.match(r"^=\d+(\.\d+)?(\+\d+(\.\d+)?)+$", v):
        return [float(x) if "." in x else int(x) for x in v[1:].split("+")]
    return None
def val(ws, a):
    return wv[ws][a].value
def inp(ws, a):
    """手入力セルの値（数式セル・空欄は None）"""
    if isf(ws, a): return None
    v = wf[ws][a].value
    if isinstance(v, str) and CONST.match(v): v = wv[ws][a].value
    return None if v in (None, "", "－", "-", "ー") else v
def num(v):
    if v is None or isinstance(v, str): return None
    return v
def fdate(v): return v.strftime("%Y/%m/%d") if isinstance(v, (datetime.date, datetime.datetime)) else ""
def ftime(v):
    if isinstance(v, datetime.datetime): v = v.time()
    if isinstance(v, datetime.time): return v.strftime("%H:%M")
    if isinstance(v, (int, float)) and 0 <= v < 2:  # Excelのシリアル時刻
        m = round((v % 1) * 1440); return f"{m//60:02d}:{m%60:02d}"
    return ""
def clean(s):
    return str(s).replace("　", " ").strip() if s is not None else ""

NIPPO = "日報（認定）"
days = []
for k in range(31):
    b = 50 * k
    dt = val(NIPPO, f"A{b+1}")
    if isinstance(dt, datetime.datetime): days.append((k, b, dt))
month = days[0][2].strftime("%Y%m")

# 0件の月でも見出しだけのCSVを出す（JUST.DB取込・分析Excelで列が必要なため）
HEADERS = {
    "T01_日報": ["日付", "今日の安全ポイント", "R-KYK重点行動目標", "汚泥含水率", "脱水汚泥_当日残量kg"],
    "T02_直別操業": ["日付", "直", "担当班", "所定時間_手入力分", "休憩時間分", "点検清掃定修分", "原水槽", "生物温度",
                  "放流濁度", "汚泥段数", "汚泥引抜_汚泥貯槽分", "減容品かさ比重", "欠員_プラオペ人数"],
    "T03_原料処理": ["日付", "直", "搬入元", "処理個数", "処理重量kg"],
    "T04_入荷明細": ["日付", "搬入元", "車両No", "重量kg", "個数"],
    "T05_減容品フレコン": ["日付", "直", "No", "重量kg"],
    "T06_廃棄物発生": ["日付", "直", "品目", "個数", "重量kg_手入力", "発生重量kg"],
    "T07_出荷明細": ["日付", "品目", "出荷先", "重量kg", "個数", "原料区分", "元データ"],
    "T08_予定外休止": ["日付", "直", "工程", "休止分"],
    "T09_設備稼働": ["日付", "時間帯", "設備", "稼働分"],
    "T10_特記事項": ["日付", "記入欄", "要因", "発生", "終了", "時間分", "内容", "備考"],
    "T11_ベール処理": ["日付", "順番", "ベール種類", "処理個数", "開始時刻", "終了時刻"],
    "T12_直送時間別": ["日付", "時間帯", "品目", "重量kg"],
}
tables = {t: [] for t in HEADERS}
def add(t, row): tables.setdefault(t, []).append(row)

SHIFTS = [("D", "I", "J", "Y"), ("N", "M", "N", "AA")]  # 直, 重量列, 個数列, 時間列
for k, b, dt in days:
    D = fdate(dt)
    r = lambda n: b + n
    # 操業実績があった日か（原料処理 or 設備稼働）
    # ---- T01 日報 ----
    add("T01_日報", {
        "日付": D,
        "今日の安全ポイント": clean(val(NIPPO, f"B{r(3)}")),
        "R-KYK重点行動目標": clean(inp(NIPPO, f"I{r(3)}")),
        "汚泥含水率": num(inp(NIPPO, f"T{r(38)}")),
        "脱水汚泥_当日残量kg": num(inp(NIPPO, f"R{r(36)}")),
    })
    # ---- T02 直別操業 ----
    for sh, wcol, ccol, tcol in SHIFTS:
        grp = inp(NIPPO, f"{'I' if sh=='D' else 'M'}{r(6)}")
        worked = any(num(val(NIPPO, f"{c}{r(30)}")) for c in ([tcol] if sh == "D" else ["AA", "AB"]))
        if not grp and not worked: continue
        add("T02_直別操業", {
            "日付": D, "直": sh, "担当班": grp or "",
            "所定時間_手入力分": num(inp(NIPPO, f"{tcol}{r(8)}")),
            "休憩時間分": num(inp(NIPPO, f"{tcol}{r(15)}")),
            "点検清掃定修分": num(inp(NIPPO, f"{tcol}{r(16)}")),
            "原水槽": num(inp(NIPPO, f"{tcol}{r(4)}")),
            "生物温度": num(inp(NIPPO, f"{tcol}{r(5)}")),
            "放流濁度": num(inp(NIPPO, f"{tcol}{r(6)}")),
            "汚泥段数": num(inp(NIPPO, f"{tcol}{r(7)}")),
            "汚泥引抜_汚泥貯槽分": num(inp(NIPPO, f"{tcol}{r(3)}")),
            "減容品かさ比重": num(inp(NIPPO, f"{tcol}{r(43)}")),
            "欠員_プラオペ人数": num(inp(NIPPO, f"AB{r(5)}")) if sh == "D" else None,
        })
    # ---- T03 原料処理（個数） ----
    for sheet, rows in ((NIPPO, (9, 10, 11)), ("日報（容リ）", (12, 13, 14)), ("日報（32条）", (15, 16))):
        for rr in rows:
            src_name = clean(val(sheet, f"C{r(rr)}"))
            for sh, cc, wc in (("D", "J", "I"), ("E", "L", "K"), ("N", "N", "M")):
                n = num(inp(sheet, f"{cc}{r(rr)}"))
                if n and src_name not in ("", "0"):
                    add("T03_原料処理", {"日付": D, "直": sh, "搬入元": src_name, "処理個数": n,
                                         "処理重量kg": num(val(sheet, f"{wc}{r(rr)}"))})
    # ---- T06 廃棄物発生 ----
    for rr in (32, 33, 34, 35):
        item = clean(val(NIPPO, f"C{r(rr)}"))
        for sh, wc, cc in (("D", "I", "J"), ("N", "M", "N")):
            n = num(inp(NIPPO, f"{cc}{r(rr)}")); w = num(inp(NIPPO, f"{wc}{r(rr)}"))
            if n or w:
                add("T06_廃棄物発生", {"日付": D, "直": sh, "品目": item, "個数": n or 0, "重量kg_手入力": w,
                                       "発生重量kg": num(val(NIPPO, f"{wc}{r(rr)}"))})
    # ---- T08 予定外休止（工程別） ----
    for rr in range(18, 29):
        proc = clean(val(NIPPO, f"X{r(rr)}"))
        for sh, tc in (("D", "Y"), ("N", "AA")):
            m = num(inp(NIPPO, f"{tc}{r(rr)}"))
            if m: add("T08_予定外休止", {"日付": D, "直": sh, "工程": proc, "休止分": m})
    # ---- T09 設備稼働時間 ----
    for rr in range(30, 39):
        eq = clean(val(NIPPO, f"X{r(rr)}")).lstrip("・")
        for sh, tc in (("D", "Y"), ("N前半(19-24時)", "AA"), ("N後半(0-7時)", "AB")):
            m = num(inp(NIPPO, f"{tc}{r(rr)}"))
            if m is not None:
                if rr == 38 and sh.startswith("N"):
                    sh = "N"  # 汚泥脱水機はN直まとめて入力
                add("T09_設備稼働", {"日付": D, "時間帯": sh, "設備": eq, "稼働分": m})
    # ---- T10 トラブル・特記事項 ----
    for rr in range(40, 51):
        for side, cols in (("D直側", ("A", "C", "D", "E", "F", "J")), ("N直側", ("L", "M", "N", "O", "P", "T"))):
            fac, st, en, mins, txt, note = (val(NIPPO, f"{c}{r(rr)}") for c in cols)
            txt = clean(txt)
            if not txt and not clean(fac): continue
            add("T10_特記事項", {"日付": D, "記入欄": side, "要因": clean(fac), "発生": ftime(st), "終了": ftime(en),
                                "時間分": round(mins) if isinstance(mins, (int, float)) and mins else None,
                                "内容": txt, "備考": clean(note)})
    # ---- T11 ベール処理記録 ----
    seq = 0
    for rr in range(7, 17):
        kind = clean(inp(NIPPO, f"AF{r(rr)}")); cnt = num(inp(NIPPO, f"AG{r(rr)}"))
        if not kind: continue
        seq += 1
        add("T11_ベール処理", {"日付": D, "順番": seq, "ベール種類": kind, "処理個数": cnt,
                               "開始時刻": ftime(val(NIPPO, f"AH{r(rr)}")), "終了時刻": ftime(val(NIPPO, f"AI{r(rr)}"))})
    # ---- T05 減容品フレコン明細 ----
    for j, sh in enumerate(("D", "E", "N")):
        col = gl(2 + 3 * k + j)
        for rr in range(11, 91):
            w = num(inp("減容品（認定）", f"{col}{rr}"))
            if w: add("T05_減容品フレコン", {"日付": D, "直": sh, "No": rr - 10, "重量kg": w})
    # ---- T04 入荷明細 ----
    ir = 5 + 50 * k
    for start, end, cname in (("AJ", "AX", "B3"), ("BC", "CE", "E3"), ("CJ", "CX", "H3"), ("DC", "DQ", "K3"), ("DV", "EJ", "N3")):
        src_name = clean(val("入荷", cname))
        for n, c in enumerate(range(ci(start), ci(end) + 1, 2), 1):
            w = num(inp("入荷", f"{gl(c)}{ir}")); q = num(inp("入荷", f"{gl(c+1)}{ir}"))
            if w: add("T04_入荷明細", {"日付": D, "搬入元": src_name, "車両No": n, "重量kg": w, "個数": q})
    for c3, name in (("T", "独自ベール（川崎・太田区）"),):
        w = num(inp("入荷", f"T{ir}")); q = num(inp("入荷", f"U{ir}"))
        if w: add("T04_入荷明細", {"日付": D, "搬入元": name, "車両No": 1, "重量kg": w, "個数": q})
    # ---- T07 出荷明細 ----（パレット出荷は認定/容リ/32条の3シートに分かれている）
    for ship_sheet, kubun in (("出荷（容リ）", "容リ"), ("出荷（32条）", "32条")):
        for L in ("B", "D", "BL"):
            w = num(val(ship_sheet, f"{L}{ir}")); q = num(val(ship_sheet, f"{gl(ci(L)+1)}{ir}"))
            if w:
                add("T07_出荷明細", {"日付": D, "品目": clean(val("出荷", "B3")) if L != "BL" else clean(val("出荷", "BL3")),
                                    "出荷先": clean(val("出荷", f"{L}4")), "重量kg": round(w), "個数": q,
                                    "原料区分": kubun, "元データ": "他ファイル連携"})
    grp = None
    for c in range(2, ci("FB") + 1):
        L = gl(c); h3 = val("出荷", f"{L}3"); h4 = val("出荷", f"{L}4")
        if h3: grp = clean(h3)
        if h4 is None or clean(h4) in ("合計", "外販小計") or grp in ("PE･PP",): continue
        w = num(val("出荷", f"{L}{ir}")); q = num(val("出荷", f"{gl(c+1)}{ir}"))
        if w:
            src_kind = "他ファイル連携" if isf("出荷", f"{L}{ir}") else "手入力"
            ws_, qs_ = const_parts("出荷", f"{L}{ir}"), const_parts("出荷", f"{gl(c+1)}{ir}")
            pairs = list(zip(ws_, qs_)) if ws_ and qs_ and len(ws_) == len(qs_) else [(round(w), q)]
            for pw, pq in pairs:   # 1セルに足し算で入っていた複数台分は1台1行に分解
                add("T07_出荷明細", {"日付": D, "品目": grp, "出荷先": clean(h4), "重量kg": pw, "個数": pq,
                                    "原料区分": "認定" if L in ("B", "D", "BL") else "", "元データ": src_kind})
    # ---- T12 直送ライン時間別重量 ----
    for item, col0 in (("PE･PP混合フラフ", 3), ("PE･PP混合減容品（直送）", ci("AK"))):
        col = gl(col0 + k)
        for rr in range(6, 30):
            w = num(val("直送", f"{col}{rr}"))
            if w:
                add("T12_直送時間別", {"日付": D, "時間帯": f"{ftime(val('直送', f'A{rr}'))}-{ftime(val('直送', f'B{rr}'))}",
                                      "品目": item, "重量kg": round(w, 1)})

# ---- 検証用：日報上の日別集計値 ----
for k, b, dt in days:
    D = fdate(dt); T = "日報（統合）"
    add("X_検証_日報集計値", {
        "日付": D,
        "原料入荷kg": val(T, f"G{b+17}"), "原料入荷個数": val(T, f"H{b+17}"),
        "原料処理kg": val(T, f"O{b+17}"), "原料処理個数": val(T, f"P{b+17}"),
        "減容品生産kg": val(NIPPO, f"O{b+22}"), "減容品生産袋数": val(NIPPO, f"P{b+22}"),
        "製品出荷kg": val(T, f"G{b+30}"), "廃棄物出荷kg": val(T, f"G{b+37}"),
        "実操業時間分": val(NIPPO, f"AC{b+10}"), "予定外休止分": val(NIPPO, f"AC{b+12}"),
    })

for t, rows in tables.items():
    keys = list(dict.fromkeys(HEADERS.get(t, []) + [k for r in rows for k in r]))
    with open(os.path.join(outdir, f"{t}_{month}.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=keys); w.writeheader()
        for r in rows: w.writerow({k: ("" if r.get(k) is None else r.get(k)) for k in keys})
    print(t, len(rows))
